from typing import List, Dict, Any
import docx


def parse_docx_document(file_path: str) -> List[Dict[str, Any]]:
    """
    Parse .docx file using python-docx.
    Returns page-like structured text list.
    """
    doc = docx.Document(file_path)
    full_text = []

    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())

    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_data:
                table_rows.append(" | ".join(row_data))
        if table_rows:
            full_text.append("\n".join(table_rows))

    text_content = "\n\n".join(full_text)

    return [{
        "page_number": 1,
        "text": text_content,
        "has_usable_text": len(text_content) > 10,
        "image_count": 0,
        "images": []
    }]
